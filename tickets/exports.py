from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TICKET_EXPORT_HEADERS = (
    'ID',
    'Titulo',
    'Fecha elaboracion',
    'Estado',
    'Prioridad',
    'Cliente',
    'Telefono cliente',
    'Asignado a',
    'Fecha creacion',
    'Ultima actualizacion',
    'Descripcion',
)


def build_ticket_export_workbook(tickets):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Tickets'

    header_fill = PatternFill(
        fill_type='solid',
        fgColor='1F2937',
    )
    header_font = Font(
        color='FFFFFF',
        bold=True,
    )

    worksheet.append(TICKET_EXPORT_HEADERS)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for ticket in tickets:
        worksheet.append(
            (
                ticket.id,
                ticket.title,
                ticket.elaboration_date.strftime('%d/%m/%Y'),
                ticket.get_status_display(),
                ticket.get_priority_display(),
                ticket.created_by.username,
                ticket.created_by.phone or '',
                ticket.assigned_to.username if ticket.assigned_to else 'Sin asignar',
                ticket.created_at.strftime('%d/%m/%Y %H:%M'),
                ticket.updated_at.strftime('%d/%m/%Y %H:%M'),
                ticket.description,
            )
        )

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(max_length + 2, 45)

    worksheet.freeze_panes = 'A2'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output
